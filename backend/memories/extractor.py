import io
import os
import re
from typing import Optional


def extract_text(file_bytes: bytes, file_type: str, url: Optional[str] = None) -> str:
    if not file_bytes and not url:
        raise ValueError("No content provided")

    file_type = (file_type or "text").lower().strip()

    if file_type == "url" or url:
        return _extract_from_url(url or file_bytes.decode("utf-8", errors="ignore").strip())
    elif file_type == "pdf":
        return _extract_from_pdf(file_bytes)
    elif file_type == "docx":
        return _extract_from_docx(file_bytes)
    elif file_type == "pptx":
        return _extract_from_pptx(file_bytes)
    elif file_type in ["image", "png", "jpg", "jpeg", "webp", "bmp", "tiff"]:
        return _extract_from_image(file_bytes)
    elif file_type == "code":
        return _extract_from_code(file_bytes)
    elif file_type in ["text", "txt", "md", "markdown"]:
        return _extract_plain_text(file_bytes)
    else:
        for extractor in [_extract_from_pdf, _extract_from_docx, _extract_plain_text]:
            try:
                result = extractor(file_bytes)
                if result and result.strip():
                    return result
            except Exception:
                continue
        raise ValueError(f"Could not extract text from file type: {file_type}")


def _extract_from_pdf(file_bytes: bytes) -> str:
    try:
        import pdfplumber
        text_parts = []
        metadata_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            if not pdf.pages:
                raise ValueError("PDF has no pages")
            if pdf.metadata:
                for key in ["Title", "Author", "Subject", "Keywords"]:
                    val = pdf.metadata.get(key, "")
                    if val and str(val).strip():
                        metadata_parts.append(f"{key}: {str(val).strip()}")
            total_pages = len(pdf.pages)
            for i, page in enumerate(pdf.pages, 1):
                try:
                    page_text = page.extract_text()
                    if page_text and page_text.strip():
                        text_parts.append(f"[Page {i}/{total_pages}]\n{page_text.strip()}")
                    tables = page.extract_tables()
                    for table in tables:
                        if table:
                            table_text = "\n".join(
                                " | ".join(str(cell or "") for cell in row)
                                for row in table if row
                            )
                            if table_text.strip():
                                text_parts.append(f"[Table on Page {i}]\n{table_text}")
                except Exception:
                    continue
        if not text_parts:
            raise ValueError("No text could be extracted from PDF")
        result_parts = []
        if metadata_parts:
            result_parts.append("Document Info:\n" + "\n".join(metadata_parts))
        result_parts.extend(text_parts)
        return "\n\n".join(result_parts).strip()
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"PDF extraction failed: {str(e)}")


def _extract_from_code(file_bytes: bytes) -> str:
    try:
        code_text = file_bytes.decode("utf-8", errors="ignore").strip()
        if not code_text:
            raise ValueError("Code file is empty")
        return code_text
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Code extraction failed: {str(e)}")


def detect_language(code: str, filename: str = "") -> str:
    ext_map = {
        ".py": "python", ".js": "javascript", ".ts": "typescript",
        ".jsx": "javascript", ".tsx": "typescript", ".java": "java",
        ".cpp": "cpp", ".c": "c", ".cs": "csharp", ".go": "go",
        ".rs": "rust", ".rb": "ruby", ".php": "php", ".swift": "swift",
        ".kt": "kotlin", ".r": "r", ".sql": "sql", ".sh": "bash",
        ".html": "html", ".css": "css", ".json": "json", ".yaml": "yaml",
        ".md": "markdown",
    }
    if filename:
        for ext, lang in ext_map.items():
            if filename.lower().endswith(ext):
                return lang
    if re.search(r"def |import |print\(|#!", code[:500]):
        return "python"
    elif re.search(r"function |const |let |var |=>", code[:500]):
        return "javascript"
    elif re.search(r"public class|System\.out|import java", code[:500]):
        return "java"
    elif re.search(r"#include|int main|cout<<", code[:500]):
        return "cpp"
    elif re.search(r"SELECT|INSERT|CREATE TABLE|FROM", code[:500], re.IGNORECASE):
        return "sql"
    return "text"


def _extract_from_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        text_parts = []
        for para in doc.paragraphs:
            if para.text and para.text.strip():
                text_parts.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        text_parts.append(cell.text.strip())
        result = "\n".join(text_parts).strip()
        if not result:
            raise ValueError("No text found in Word document")
        return result
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Word document extraction failed: {str(e)}")


def _extract_from_pptx(file_bytes: bytes) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(file_bytes))
        text_parts = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text and shape.text.strip():
                    slide_texts.append(shape.text.strip())
            if slide_texts:
                text_parts.append(f"Slide {slide_num}:\n" + "\n".join(slide_texts))
        result = "\n\n".join(text_parts).strip()
        if not result:
            raise ValueError("No text found in PowerPoint")
        return result
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"PowerPoint extraction failed: {str(e)}")


def _extract_from_image(file_bytes: bytes) -> str:
    IS_PRODUCTION = bool(os.getenv("RENDER", False))

    if IS_PRODUCTION:
        try:
            from PIL import Image
            image = Image.open(io.BytesIO(file_bytes))
            width, height = image.size
            mode = image.mode
            return (
                f"Image uploaded successfully.\n"
                f"Dimensions: {width}x{height}\n"
                f"Mode: {mode}\n"
                f"Note: OCR text extraction is not available in production. "
                f"Please describe the image content in the title field."
            )
        except Exception:
            return "Image uploaded. Please describe the content in the title field."

    try:
        import pytesseract
        from PIL import Image

        tesseract_paths = [
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
            "/usr/bin/tesseract",
            "/usr/local/bin/tesseract",
        ]
        for path in tesseract_paths:
            if os.path.exists(path):
                pytesseract.pytesseract.tesseract_cmd = path
                break

        image = Image.open(io.BytesIO(file_bytes))
        if image.mode not in ("RGB", "L"):
            image = image.convert("RGB")

        text = pytesseract.image_to_string(image, config="--psm 3")
        result = text.strip()
        if not result:
            raise ValueError("No text found in image")
        return result
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Image OCR failed: {str(e)}")


def _extract_from_url(url: str) -> str:
    if not url or not url.strip():
        raise ValueError("URL cannot be empty")
    url = url.strip()
    if not url.startswith(("http://", "https://")):
        url = "https://" + url

    # Block YouTube URLs since that feature is removed
    if _is_youtube_url(url):
        raise ValueError(
            "YouTube URLs are not supported. "
            "Please paste the article or webpage URL directly."
        )

    try:
        import trafilatura
        downloaded = trafilatura.fetch_url(url)
        if not downloaded:
            raise ValueError("Could not fetch URL — check if it is accessible")
        text = trafilatura.extract(
            downloaded,
            include_comments=False,
            include_tables=True,
            no_fallback=False,
        )
        if not text or not text.strip():
            raise ValueError("No readable text found at this URL")
        return text.strip()
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"URL extraction failed: {str(e)}")


def _extract_plain_text(file_bytes: bytes) -> str:
    for encoding in ["utf-8", "latin-1", "cp1252", "ascii"]:
        try:
            text = file_bytes.decode(encoding).strip()
            if text:
                return text
        except (UnicodeDecodeError, ValueError):
            continue
    raise ValueError("Could not decode file as plain text")


def _is_youtube_url(url: str) -> bool:
    youtube_patterns = [
        r"youtube\.com",
        r"youtu\.be",
    ]
    return any(re.search(p, url) for p in youtube_patterns)